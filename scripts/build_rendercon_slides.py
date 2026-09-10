"""Build RenderCon Proxy vs BFF slides using Slidesgo Tech Startup palette."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "docs" / "RenderCon_Proxy_vs_BFF_Slides.pptx"

# Slidesgo Tech Startup palette
NAVY = RGBColor(0x07, 0x37, 0x63)
BLUE = RGBColor(0x0B, 0x53, 0x94)
TEAL = RGBColor(0x00, 0xD6, 0xC0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xEF, 0xEF, 0xEF)
GRAY = RGBColor(0xCC, 0xCC, 0xCC)
DARK = RGBColor(0x42, 0x38, 0x64)
MUTED = RGBColor(0x6B, 0x72, 0x80)
ROW_A = RGBColor(0xF5, 0xF7, 0xFA)
ROW_B = RGBColor(0xE8, 0xF8, 0xF6)


def set_run(run, text: str, size: int, bold: bool = False, color: RGBColor = WHITE):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def add_rect(slide, left, top, width, height, fill: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    lines: list[tuple[str, int, bool, RGBColor]],
    align=PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for text, size, bold, color in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        run = p.add_run()
        set_run(run, text, size, bold, color)
        p.space_after = Pt(6)
    return box


def blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def paint_bg(slide, color: RGBColor = NAVY):
    add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(7.5), color)


def accent_bar(slide, color: RGBColor = TEAL):
    add_rect(slide, Inches(0), Inches(0), Inches(0.18), Inches(7.5), color)


def footer(slide, page: str):
    add_textbox(
        slide,
        Inches(0.5),
        Inches(7.05),
        Inches(10),
        Inches(0.35),
        [("RenderCon · Proxy vs BFF · GreenKey Realty", 11, False, GRAY)],
    )
    add_textbox(
        slide,
        Inches(11.5),
        Inches(7.05),
        Inches(1.5),
        Inches(0.35),
        [(page, 11, False, GRAY)],
        align=PP_ALIGN.RIGHT,
    )


def title_slide(prs):
    s = blank_slide(prs)
    paint_bg(s, NAVY)
    add_rect(s, Inches(0), Inches(5.8), Inches(13.333), Inches(1.7), BLUE)
    add_textbox(
        s,
        Inches(0.7),
        Inches(1.6),
        Inches(12),
        Inches(0.5),
        [("RENDERCON TECH TALK", 14, True, TEAL)],
    )
    add_textbox(
        s,
        Inches(0.7),
        Inches(2.2),
        Inches(12),
        Inches(1.2),
        [("PROXY vs BFF", 48, True, WHITE)],
    )
    add_textbox(
        s,
        Inches(0.7),
        Inches(3.4),
        Inches(12),
        Inches(1.2),
        [
            ("Building a Backend for Frontend with Next.js + NestJS", 22, False, TEAL),
            ("GreenKey Realty · Nyeri real estate product", 16, False, LIGHT),
        ],
    )
    add_textbox(
        s,
        Inches(0.7),
        Inches(6.15),
        Inches(12),
        Inches(0.9),
        [
            ("[Your Name]  ·  fullstack engineer", 16, False, WHITE),
            ("30–40 minutes  ·  talk + live demo", 14, False, GRAY),
        ],
    )


def agenda_slide(prs):
    s = blank_slide(prs)
    paint_bg(s, WHITE)
    accent_bar(s)
    add_textbox(s, Inches(0.6), Inches(0.35), Inches(10), Inches(0.6), [("AGENDA", 32, True, NAVY)])
    items = [
        ("01", "Hook — receptionist vs assistant"),
        ("02", "Why many backends break “Next-only” apps"),
        ("03", "Proxy vs BFF definitions"),
        ("04", "GreenKey architecture & ownership"),
        ("05", "Product stories (auth, save, M-Pesa, dual-write)"),
        ("06", "Live demo + Swagger as mobile view"),
        ("07", "Takeaways & Q&A"),
    ]
    y = 1.2
    for num, text in items:
        add_rect(s, Inches(0.6), Inches(y), Inches(0.7), Inches(0.55), TEAL)
        add_textbox(s, Inches(0.7), Inches(y + 0.05), Inches(0.55), Inches(0.45), [(num, 14, True, NAVY)])
        add_textbox(s, Inches(1.55), Inches(y + 0.05), Inches(10), Inches(0.45), [(text, 18, False, DARK)])
        y += 0.7
    footer(s, "2")


def quote_slide(prs):
    s = blank_slide(prs)
    paint_bg(s, NAVY)
    add_textbox(
        s,
        Inches(1),
        Inches(1.8),
        Inches(11),
        Inches(3.5),
        [
            ("Proxy = the receptionist who points you", 26, False, WHITE),
            ("to the right department.", 26, False, WHITE),
            ("", 12, False, WHITE),
            ("BFF = the assistant who visits three departments,", 26, False, TEAL),
            ("collects the paperwork, staples it, and hands", 26, False, TEAL),
            ("you a one-page summary.", 26, False, TEAL),
        ],
    )
    add_textbox(
        s,
        Inches(1),
        Inches(5.8),
        Inches(11),
        Inches(0.6),
        [("— Spine of this talk", 16, False, GRAY)],
    )
    footer(s, "3")


def section_slide(prs, number: str, title: str, subtitle: str, page: str):
    s = blank_slide(prs)
    paint_bg(s, BLUE)
    add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(0.25), TEAL)
    add_textbox(s, Inches(0.8), Inches(2.4), Inches(11), Inches(0.5), [(number, 18, True, TEAL)])
    add_textbox(s, Inches(0.8), Inches(3.0), Inches(11), Inches(1), [(title, 40, True, WHITE)])
    add_textbox(s, Inches(0.8), Inches(4.2), Inches(11), Inches(0.8), [(subtitle, 18, False, LIGHT)])
    footer(s, page)


def bullets_slide(prs, title: str, bullets: list[str], page: str, intro: str | None = None):
    s = blank_slide(prs)
    paint_bg(s, WHITE)
    accent_bar(s)
    add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.7), [(title, 28, True, NAVY)])
    y = 1.2
    if intro:
        add_textbox(s, Inches(0.6), Inches(y), Inches(12), Inches(0.6), [(intro, 16, False, DARK)])
        y = 1.9
    for b in bullets:
        add_rect(s, Inches(0.65), Inches(y + 0.12), Inches(0.18), Inches(0.18), TEAL)
        add_textbox(s, Inches(1.05), Inches(y), Inches(11.5), Inches(0.55), [(b, 17, False, DARK)])
        y += 0.65
    footer(s, page)


def two_col_slide(
    prs,
    title: str,
    left_title: str,
    left_items: list[str],
    right_title: str,
    right_items: list[str],
    page: str,
):
    s = blank_slide(prs)
    paint_bg(s, WHITE)
    accent_bar(s)
    add_textbox(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.6), [(title, 28, True, NAVY)])
    add_rect(s, Inches(0.5), Inches(1.1), Inches(5.9), Inches(5.4), NAVY)
    add_textbox(s, Inches(0.8), Inches(1.35), Inches(5.3), Inches(0.5), [(left_title, 20, True, TEAL)])
    y = 2.1
    for item in left_items:
        add_textbox(s, Inches(0.8), Inches(y), Inches(5.3), Inches(0.55), [(f"•  {item}", 14, False, WHITE)])
        y += 0.6
    add_rect(s, Inches(6.8), Inches(1.1), Inches(5.9), Inches(5.4), BLUE)
    add_textbox(s, Inches(7.1), Inches(1.35), Inches(5.3), Inches(0.5), [(right_title, 20, True, TEAL)])
    y = 2.1
    for item in right_items:
        add_textbox(s, Inches(7.1), Inches(y), Inches(5.3), Inches(0.55), [(f"•  {item}", 14, False, WHITE)])
        y += 0.6
    footer(s, page)


def architecture_slide(prs):
    s = blank_slide(prs)
    paint_bg(s, WHITE)
    accent_bar(s)
    add_textbox(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.55), [("ARCHITECTURE", 28, True, NAVY)])
    add_textbox(
        s,
        Inches(0.6),
        Inches(0.95),
        Inches(12),
        Inches(0.45),
        [("Two-layer BFF: Next is the web edge · Nest /v1 is the domain BFF", 15, False, DARK)],
    )
    boxes = [
        (0.5, "Browser", "UI only", NAVY, WHITE),
        (3.1, "Next.js", "cookies · /api · proxy", NAVY, WHITE),
        (5.7, "Nest /v1", "OTP · billing · writes", TEAL, NAVY),
        (8.3, "Data plane", "PG · Redis · Sanity", NAVY, WHITE),
        (10.6, "Providers", "M-Pesa · SMS · email", BLUE, WHITE),
    ]
    for x, title, sub, bg, fg in boxes:
        add_rect(s, Inches(x), Inches(2.0), Inches(2.3), Inches(1.8), bg)
        add_textbox(
            s,
            Inches(x + 0.12),
            Inches(2.35),
            Inches(2.05),
            Inches(1.2),
            [(title, 15, True, fg), (sub, 11, False, fg)],
            align=PP_ALIGN.CENTER,
        )
    add_textbox(
        s,
        Inches(0.6),
        Inches(4.2),
        Inches(12),
        Inches(2.2),
        [
            ("Flow", 16, True, NAVY),
            ("Browser → Next (httpOnly cookies + Route Handlers) → Nest BFF", 15, False, DARK),
            ("Nest talks to Postgres, Redis, Sanity write API, Daraja, AT/Infobip", 15, False, DARK),
            ("Next may GROQ-read Sanity for public pages; secrets never leave Nest", 15, False, DARK),
            ("Mobile uses the same /v1 with Bearer tokens (Swagger demo)", 15, False, DARK),
        ],
    )
    footer(s, "8")


def ownership_slide(prs):
    s = blank_slide(prs)
    paint_bg(s, WHITE)
    accent_bar(s)
    add_textbox(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.55), [("DATA OWNERSHIP", 28, True, NAVY)])
    rows = [
        ("Concern", "Owner", True),
        ("Listing content & media", "Sanity", False),
        ("Users, JWT refresh, roles, subscriptions", "Postgres + Nest", False),
        ("OTP codes, rate limits, payment locks", "Redis", False),
        ("M-Pesa / SMS / WhatsApp / email secrets", "Nest only", False),
        ("PropertyMirror / Lead ops rows", "Postgres via Nest", False),
        ("Agent dashboard listing UI", "Sanity (writes via Nest)", False),
    ]
    y = 1.15
    for i, (a, b, header) in enumerate(rows):
        left_bg = NAVY if header else (ROW_A if i % 2 == 0 else WHITE)
        right_bg = TEAL if header else (ROW_B if i % 2 == 0 else WHITE)
        left_fg = WHITE if header else DARK
        right_fg = NAVY if header else DARK
        add_rect(s, Inches(0.5), Inches(y), Inches(7.2), Inches(0.7), left_bg)
        add_rect(s, Inches(7.7), Inches(y), Inches(4.8), Inches(0.7), right_bg)
        add_textbox(s, Inches(0.7), Inches(y + 0.15), Inches(6.8), Inches(0.45), [(a, 14, header, left_fg)])
        add_textbox(s, Inches(7.9), Inches(y + 0.15), Inches(4.4), Inches(0.45), [(b, 14, header, right_fg)])
        y += 0.72
    footer(s, "9")


def stories_slide(prs):
    s = blank_slide(prs)
    paint_bg(s, WHITE)
    accent_bar(s)
    add_textbox(
        s,
        Inches(0.6),
        Inches(0.3),
        Inches(12),
        Inches(0.55),
        [("GREENKEY PRODUCT STORIES", 26, True, NAVY)],
    )
    cards = [
        ("Auth", "OTP in Nest + Redis\nWeb: httpOnly cookies\nMobile: Bearer tokens"),
        ("Aggregation", "Saved listings =\nPostgres IDs +\nSanity property cards"),
        ("Billing", "M-Pesa STK + callback\nstay on Nest\nNo Daraja keys in browser"),
        ("Dual-write", "Create listing →\nSanity doc +\nPropertyMirror"),
    ]
    x = 0.45
    for title, body in cards:
        add_rect(s, Inches(x), Inches(1.2), Inches(3.0), Inches(5.0), NAVY)
        add_rect(s, Inches(x), Inches(1.2), Inches(3.0), Inches(0.7), TEAL)
        add_textbox(
            s,
            Inches(x + 0.15),
            Inches(1.3),
            Inches(2.7),
            Inches(0.5),
            [(title, 16, True, NAVY)],
            align=PP_ALIGN.CENTER,
        )
        lines = [(ln, 13, False, WHITE) for ln in body.split("\n")]
        add_textbox(s, Inches(x + 0.2), Inches(2.2), Inches(2.6), Inches(3.6), lines)
        x += 3.2
    footer(s, "10")


def screen_map_slide(prs):
    s = blank_slide(prs)
    paint_bg(s, WHITE)
    accent_bar(s)
    add_textbox(
        s,
        Inches(0.6),
        Inches(0.3),
        Inches(12),
        Inches(0.55),
        [("ONE BACKEND · MANY VIEWS", 26, True, NAVY)],
    )
    rows = [
        ("Buyer web", "/properties", "Sanity read + optional session"),
        ("Auth", "/sign-in", "Next /api/auth → Nest OTP"),
        ("Buyer account", "/saved", "Postgres + Sanity join"),
        ("Billing", "/pricing", "Nest plan + STK / bank"),
        ("Agent", "/dashboard", "Same BFF + plan gate"),
        ("Admin", "/admin/payments", "Ops slice of /v1"),
        ("CMS", "/studio", "Raw Sanity documents"),
        ("Mobile view", ":4000/docs", "Swagger + Bearer tokens"),
    ]
    y = 1.05
    add_rect(s, Inches(0.5), Inches(y), Inches(12.3), Inches(0.45), NAVY)
    add_textbox(s, Inches(0.7), Inches(y + 0.05), Inches(3), Inches(0.35), [("View", 12, True, WHITE)])
    add_textbox(s, Inches(4.0), Inches(y + 0.05), Inches(3.5), Inches(0.35), [("Surface", 12, True, WHITE)])
    add_textbox(s, Inches(7.8), Inches(y + 0.05), Inches(4.8), Inches(0.35), [("What it shows", 12, True, WHITE)])
    y += 0.5
    for i, (a, b, c) in enumerate(rows):
        bg = ROW_A if i % 2 == 0 else WHITE
        add_rect(s, Inches(0.5), Inches(y), Inches(12.3), Inches(0.55), bg)
        add_textbox(s, Inches(0.7), Inches(y + 0.1), Inches(3), Inches(0.35), [(a, 13, True, DARK)])
        add_textbox(s, Inches(4.0), Inches(y + 0.1), Inches(3.5), Inches(0.35), [(b, 13, False, BLUE)])
        add_textbox(s, Inches(7.8), Inches(y + 0.1), Inches(4.8), Inches(0.35), [(c, 13, False, DARK)])
        y += 0.55
    footer(s, "11")


def demo_slide(prs):
    s = blank_slide(prs)
    paint_bg(s, NAVY)
    add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(0.25), TEAL)
    add_textbox(s, Inches(0.7), Inches(0.5), Inches(12), Inches(0.6), [("LIVE DEMO", 32, True, TEAL)])
    steps = [
        "1. Browse / → /properties",
        "2. OTP sign-in → cookies (not raw JWTs in JSON)",
        "3. Save a listing (aggregation)",
        "4. /pricing → simulate STK → agent unlock",
        "5. Create listing → open /studio (dual-write)",
        "6. Incognito: Contact agent → /dashboard/leads",
        "7. Swagger /docs = mobile Bearer view",
    ]
    y = 1.4
    for step in steps:
        add_textbox(s, Inches(0.9), Inches(y), Inches(11.5), Inches(0.5), [(step, 18, False, WHITE)])
        y += 0.65
    footer(s, "12")


def checklist_slide(prs):
    s = blank_slide(prs)
    paint_bg(s, WHITE)
    accent_bar(s)
    add_textbox(
        s,
        Inches(0.6),
        Inches(0.3),
        Inches(12),
        Inches(0.55),
        [("DEMO PREP — 8 TABS", 26, True, NAVY)],
    )
    left = [
        "http://localhost:3000",
        "/sign-in",
        "/saved",
        "/pricing",
    ]
    right = [
        "/dashboard",
        "/dashboard/admin/payments",
        "/studio",
        "http://localhost:4000/docs",
    ]
    add_textbox(s, Inches(0.6), Inches(1.2), Inches(5.5), Inches(0.4), [("Web surfaces", 16, True, BLUE)])
    add_textbox(s, Inches(7.0), Inches(1.2), Inches(5.5), Inches(0.4), [("Ops / mobile", 16, True, BLUE)])
    y = 1.8
    for a, b in zip(left, right):
        add_textbox(s, Inches(0.6), Inches(y), Inches(5.8), Inches(0.5), [(f"•  {a}", 16, False, DARK)])
        add_textbox(s, Inches(7.0), Inches(y), Inches(5.8), Inches(0.5), [(f"•  {b}", 16, False, DARK)])
        y += 0.7
    add_textbox(
        s,
        Inches(0.6),
        Inches(5.2),
        Inches(12),
        Inches(1.2),
        [
            ("Run: pnpm docker:up · pnpm dev:api · pnpm/npm run dev", 14, False, NAVY),
            ("Demo-safe: console OTP · M-Pesa simulate-complete", 14, False, DARK),
        ],
    )
    footer(s, "13")


def rule_slide(prs):
    s = blank_slide(prs)
    paint_bg(s, BLUE)
    add_textbox(
        s,
        Inches(0.8),
        Inches(2.0),
        Inches(11.5),
        Inches(3.5),
        [
            ("RULE OF THUMB", 18, True, TEAL),
            ("", 10, False, WHITE),
            ("If you’re only forwarding,", 28, False, WHITE),
            ("it’s a proxy.", 28, False, WHITE),
            ("", 12, False, WHITE),
            ("If you’re shaping for a client", 28, False, TEAL),
            ("and protecting secrets,", 28, False, TEAL),
            ("it’s a BFF.", 28, False, TEAL),
        ],
    )
    footer(s, "14")


def takeaways_slide(prs):
    s = blank_slide(prs)
    paint_bg(s, WHITE)
    accent_bar(s)
    add_textbox(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.55), [("TAKEAWAYS", 28, True, NAVY)])
    items = [
        "Proxy directs; BFF delivers a UI-shaped answer.",
        "Don’t put provider secrets in the client.",
        "One /v1 contract serves web + mobile (cookies vs Bearer).",
        "CMS ≠ system of record for accounts and payments.",
        "Why not M-Pesa in Next? Callbacks, secrets, mobile, ops.",
    ]
    y = 1.2
    for i, item in enumerate(items, 1):
        add_rect(s, Inches(0.6), Inches(y), Inches(12.1), Inches(0.85), NAVY if i % 2 else BLUE)
        add_textbox(
            s,
            Inches(0.9),
            Inches(y + 0.2),
            Inches(11.5),
            Inches(0.5),
            [(f"{i}.  {item}", 16, False, WHITE)],
        )
        y += 1.0
    footer(s, "15")


def qa_slide(prs):
    s = blank_slide(prs)
    paint_bg(s, NAVY)
    add_rect(s, Inches(0), Inches(5.9), Inches(13.333), Inches(1.6), TEAL)
    add_textbox(s, Inches(0.8), Inches(2.2), Inches(11.5), Inches(1.2), [("QUESTIONS?", 48, True, WHITE)])
    add_textbox(
        s,
        Inches(0.8),
        Inches(3.5),
        Inches(11.5),
        Inches(1.5),
        [
            ("[Your Name]  ·  GreenKey Realty", 20, False, LIGHT),
            ("Docs: TECH_TALK_BFF.md · MOBILE_API.md · RENDERCON_FULL_TALK_SCRIPT.md", 14, False, GRAY),
            ("Next.js BFF guide · classic BFF talk (YouTube)", 14, False, GRAY),
        ],
    )
    add_textbox(
        s,
        Inches(0.8),
        Inches(6.25),
        Inches(11.5),
        Inches(0.8),
        [("Thanks — let’s discuss architecture choices", 18, True, NAVY)],
    )
    footer(s, "16")


def thanks_slide(prs):
    s = blank_slide(prs)
    paint_bg(s, WHITE)
    accent_bar(s, NAVY)
    add_textbox(s, Inches(0.7), Inches(2.2), Inches(12), Inches(1), [("THANKS", 44, True, NAVY)])
    add_textbox(
        s,
        Inches(0.7),
        Inches(3.4),
        Inches(12),
        Inches(2.5),
        [
            ("Presentation visuals adapted from a Slidesgo Tech Startup template.", 16, False, DARK),
            ("Please keep this attribution slide for license compliance.", 14, False, MUTED),
            ("Palette inspired by Squada One / Roboto Condensed theme colors.", 14, False, MUTED),
            ("slidesgo.com", 14, False, BLUE),
        ],
    )
    footer(s, "17")


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    title_slide(prs)
    agenda_slide(prs)
    quote_slide(prs)
    section_slide(prs, "01", "WHY THIS MATTERS", "Many backends. One UI. Secrets must stay server-side.", "4")
    bullets_slide(
        prs,
        "THE PROBLEM",
        [
            "Modern React apps need identity, CMS, payments, and messaging.",
            "Putting all of that in the browser (or only in Next) leaks secrets.",
            "You also duplicate logic when mobile arrives.",
            "GreenKey: Sanity + Postgres + Redis + M-Pesa + SMS/WhatsApp/email.",
        ],
        "5",
        intro="If every client talks to every provider, coupling explodes.",
    )
    two_col_slide(
        prs,
        "PROXY vs BFF",
        "PROXY (receptionist)",
        [
            "Route / gate / rewrite",
            "“Go to Billing” / redirect to sign-in",
            "Next.js proxy.ts",
            "Protects /dashboard, refreshes cookies",
            "Does not assemble business data",
        ],
        "BFF (assistant)",
        [
            "Shape data for this UI",
            "Call PG + Sanity + M-Pesa → one response",
            "Nest /v1 + Next /api Route Handlers",
            "Hides Daraja & Sanity write tokens",
            "Shared contract for web + mobile",
        ],
        "6",
    )
    section_slide(
        prs,
        "02",
        "SOLUTION",
        "Next edge + Nest domain BFF — matches Next docs, ready for mobile.",
        "7",
    )
    architecture_slide(prs)
    ownership_slide(prs)
    stories_slide(prs)
    screen_map_slide(prs)
    demo_slide(prs)
    checklist_slide(prs)
    rule_slide(prs)
    takeaways_slide(prs)
    qa_slide(prs)
    thanks_slide(prs)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
